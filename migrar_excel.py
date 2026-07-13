"""
migrar_excel.py — Genera/actualiza datos.json a partir de la hoja 'AFIP'
del Excel original. Solo necesario si querés re-migrar desde el Excel;
la app ya viene con datos.json precargado.

Uso: python migrar_excel.py AFIP-Para_Claude.xlsx
"""
import sys, json
import openpyxl

BLOQUES = [('AL','AM'), ('AF','AG'), ('Z','AA'), ('T','U'), ('N','O'), ('H','I'), ('B','C')]
INF = {
    'G': ['2026-01','2026-02','2026-03','2026-04','2026-05','2026-06'],
    'M': ['2025-07','2025-08','2025-09','2025-10','2025-11','2025-12'],
    'S': ['2025-01','2025-02','2025-03','2025-04','2025-05','2025-06'],
    'Y': ['2024-07','2024-08','2024-09','2024-10','2024-11','2024-12'],
}

def main(path_xlsx, path_json='datos.json'):
    ws = openpyxl.load_workbook(path_xlsx, data_only=True)['AFIP']
    serie, infl = {}, {}
    for cf, cm in BLOQUES:
        for r in range(2, 14):
            f, m = ws[f'{cf}{r}'].value, ws[f'{cm}{r}'].value
            if f is None:
                continue
            k = f'{f.year}-{f.month:02d}'
            if m is not None:
                serie[k] = float(m)
            elif k not in serie:
                serie[k] = 0.0
    for col, meses in INF.items():
        for i, mes in enumerate(meses):
            v = ws[f'{col}{8+i}'].value
            if v is not None:
                infl[mes] = float(v)
    try:
        base = json.load(open(path_json, encoding='utf-8'))
    except FileNotFoundError:
        base = {"facturacion": {}, "inflacion": {}, "inflacion_estimada": {},
                "escalas": {}, "categoria_meta": "G", "contribuyente": "", "cuit": ""}
    base['facturacion'].update(serie)
    for k, v in infl.items():
        base.setdefault('inflacion_estimada', {}).setdefault(k, v)
    json.dump(base, open(path_json, 'w', encoding='utf-8'), indent=2, ensure_ascii=False)
    print(f"OK: {len(serie)} meses de facturación, {len(infl)} tasas → {path_json}")

if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else 'AFIP-Para_Claude.xlsx')
